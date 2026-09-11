import os
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = 'docs/apnacare_product_prd.docx'

def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    props.append(shd)

def table_borders(table):
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = OxmlElement('w:' + edge)
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), '4')
        node.set(qn('w:color'), 'D9D9D9')
        borders.append(node)
    props.append(borders)

def set_cell(cell, value, bold=False, color='222222'):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(str(value))
    r.bold = bold
    r.font.name = 'Aptos'
    r.font.size = Pt(9.2)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t)
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, True, 'FFFFFF')
        shade(t.rows[0].cells[i], '2F6B3B')
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
            if n % 2:
                shade(cells[i], 'F1F6F0')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)

def numbered(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)

os.makedirs('docs', exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].font.color.rgb = RGBColor(0x22, 0x22, 0x22)
styles['Normal'].paragraph_format.space_after = Pt(7)
styles['Normal'].paragraph_format.line_spacing = 1.12
for name, size in [('Title', 28), ('Heading 1', 18), ('Heading 2', 13), ('Heading 3', 11)]:
    styles[name].font.name = 'Aptos Display' if name in ('Title', 'Heading 1') else 'Aptos'
    styles[name]._element.rPr.rFonts.set(qn('w:ascii'), styles[name].font.name)
    styles[name]._element.rPr.rFonts.set(qn('w:hAnsi'), styles[name].font.name)
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

p = doc.add_paragraph(style='Title')
p.add_run('ApnaCare Product Requirements Document')
p = doc.add_paragraph()
r = p.add_run('Emergency to everyday healthcare platform')
r.bold = True
r.font.size = Pt(15)
r.font.color.rgb = RGBColor(0x2F, 0x6B, 0x3B)
doc.add_paragraph('Company strategy, product requirements, access model, and frontend implementation blueprint')
add_table(doc, ['Document field', 'Decision'], [
    ('Product', 'ApnaCare'),
    ('Company type', 'Multi-sided healthcare network and provider operating platform'),
    ('Launch wedge', 'Emergency Care Room and ambulance coordination'),
    ('Initial build', 'Flutter frontend with realistic dummy repositories'),
    ('Audience', 'Founder, product, engineering, provider partners, and investors'),
])
doc.add_paragraph('This document defines what ApnaCare is building, why it should exist as a company, how patients and providers benefit, how access is controlled, and how the current frontend prototype should evolve. The first release proves one complete emergency workflow; later releases turn the same network into the everyday care and provider operations layer.')

h(doc, 'Executive decision')
doc.add_paragraph('ApnaCare will be an emergency-to-everyday healthcare platform. Patients and families use it for emergency access, care coordination, appointments, records, bills, and future medicine services. Hospitals, clinics, independent doctors, ambulance providers, and diagnostic partners use the same network to run their workflows and reach patients.')
doc.add_paragraph('ApnaCare owns the platform, permission primitives, emergency network, and marketplace. Each provider organization receives an owner account with full control of its own tenant. The owner can invite staff, create custom roles from ApnaCare permissions, configure departments and services, publish availability, and manage billing without waiting for ApnaCare support.')

h(doc, '1 Company thesis and goals')
h(doc, '1.1 Problem')
doc.add_paragraph('Healthcare is fragmented when coordination matters most. A family may call an ambulance, search for a hospital, carry paper reports, phone a doctor, wait at reception, and pay through disconnected systems. Providers operate in separate appointment calendars, billing tools, spreadsheets, phone lines, and manually maintained records.')
doc.add_paragraph('This creates lost time, poor handoffs, empty appointment capacity, no-shows, slow billing, missing emergency information, and no shared accountability across a care journey.')
h(doc, '1.2 Company goals', 2)
for item in [
    'Make emergency care coordination faster, more informed, and auditable.',
    'Give every participating provider a useful operating system before a full network exists.',
    'Create a trusted patient identity, emergency card, and consent layer across providers.',
    'Create a network effect: patient demand attracts providers; provider readiness makes the patient app more valuable.',
    'Build sustainable revenue through provider subscriptions, workflow modules, network fees, and family services.',
]: bullet(doc, item)
h(doc, '1.3 Success definition', 2)
add_table(doc, ['Stage', 'Evidence of success'], [
    ('Frontend demo', 'One emergency incident can be followed across patient, dispatcher, ambulance, hospital, doctor, billing, and admin workspaces.'),
    ('Pilot', 'A provider owner configures staff, roles, services, schedules, billing, and emergency readiness without engineering help.'),
    ('Network', 'Patients use ApnaCare for emergency readiness and appointments; providers receive demand and operational visibility.'),
    ('Company', 'ApnaCare becomes the coordination layer providers and families do not want to operate without.'),
])

h(doc, '2 Product positioning')
doc.add_paragraph('ApnaCare is the emergency-to-everyday care layer for families and the operating system for healthcare providers. The reason to install is emergency readiness: a family can prepare an emergency card, contacts, location, and trusted providers before a crisis. The reason to keep using it is appointments, records, bills, reminders, referrals, and medicine services.')
h(doc, '2.1 Product surfaces', 2)
add_table(doc, ['Surface', 'Primary job', 'Business value'], [
    ('Patient app', 'Request help, find care, manage family health, track events', 'Creates demand and daily habit'),
    ('Ambulance workspace', 'Accept, dispatch, navigate, communicate, close incidents', 'Turns the wedge into a network'),
    ('Doctor workspace', 'See queue, context, notes, prescriptions, follow-ups', 'Makes clinicians active participants'),
    ('Provider workspace', 'Run appointments, queues, staff, services, billing, reports', 'Creates recurring SaaS value'),
    ('ApnaCare control plane', 'Verify providers, manage safety, plans, disputes, network', 'Protects trust and monetization'),
])
h(doc, '2.2 Product principles', 2)
for item in [
    'Emergency-first, but not emergency-only.',
    'Configuration over customization: provider owners configure workflows without code.',
    'Permission-aware by default: every person sees only what their task requires.',
    'Explain recommendations using readiness, capability, cost, distance, and availability.',
    'Every important action has a status, owner, timestamp, and audit trail.',
    'AI assists navigation and organization; it does not diagnose or prescribe.',
]: bullet(doc, item)

h(doc, '3 Customers and participants')
add_table(doc, ['Participant', 'Pain point', 'ApnaCare value'], [
    ('Patient or caregiver', 'Unclear next step; repeated information', 'Emergency card, care search, booking, records, reminders'),
    ('Ambulance provider', 'Phone-based dispatch and weak hospital coordination', 'Live requests, capability matching, navigation, handoff, settlement'),
    ('Hospital or clinic owner', 'Disconnected tools and manual administration', 'Configurable operations, demand, workflows, revenue visibility'),
    ('Doctor', 'Unclear queue and missing patient context', 'Today view, reports, notes, prescriptions, availability'),
    ('Operations and billing staff', 'Repeated entry and unclear responsibility', 'Role-specific queues, invoices, receipts, refunds, auditability'),
    ('ApnaCare team', 'Trust is difficult across a distributed network', 'Verification, policy, SLA, plans, and dispute controls'),
])

h(doc, '4 Organization and access model')
doc.add_paragraph('Provider type is not a role. A hospital, clinic, ambulance company, diagnostic centre, pharmacy, or independent practice is an organization. A doctor, dispatcher, driver, receptionist, billing operator, or administrator is a person’s responsibility inside an organization. One person can belong to several organizations with different roles and scopes.')
h(doc, '4.1 Core entities', 2)
add_table(doc, ['Entity', 'Purpose'], [
    ('Person', 'One identity for login, profile, security history, and notifications'),
    ('Organization', 'Hospital, clinic, ambulance provider, diagnostic centre, pharmacy, or ApnaCare'),
    ('Branch and department', 'Operational boundaries inside an organization'),
    ('Membership', 'Connects a person to an organization and its status'),
    ('Role', 'Named bundle of platform or organization permissions'),
    ('Permission', 'Atomic action such as appointments.view or billing.refund'),
    ('Scope', 'Own, assigned, department, branch, organization, network, or platform'),
    ('Policy and audit event', 'Contextual rule and immutable sensitive-action history'),
])
h(doc, '4.2 Owner and custom role workflow', 2)
for item in [
    'ApnaCare verifies and creates the provider organization.',
    'The organization owner receives full control of that organization only.',
    'The owner configures branches, departments, services, fees, schedules, staff, and billing.',
    'The owner creates roles from ApnaCare permission primitives or recommended templates.',
    'The owner invites staff and assigns roles with scope, branch, department, and optional expiry.',
    'The owner reviews access, revokes sessions, and inspects audit events independently.',
]: numbered(doc, item)
h(doc, '4.3 Example base roles', 2)
add_table(doc, ['Organization', 'Base roles', 'Owner-created examples'], [
    ('Hospital', 'Owner, admin, doctor, nurse, receptionist, billing, dispatcher', 'Night coordinator, ICU billing reviewer, branch quality lead'),
    ('Clinic', 'Owner, doctor, receptionist, billing', 'Visiting consultant, part-time front desk'),
    ('Ambulance provider', 'Owner, dispatcher, driver, medical crew', 'Fleet supervisor, night dispatcher'),
    ('Independent doctor', 'Owner-doctor, assistant, billing', 'Teleconsult coordinator, hospital liaison'),
])
h(doc, '4.4 Permission families', 2)
for item in [
    'appointments.view, create, reschedule, cancel, check_in',
    'patients.view, edit, assign, records.view, records.upload, records.share',
    'billing.view, invoice.create, payment.record, refund, export',
    'ambulance.request.view, accept, dispatch, assign, location.view, close',
    'schedules.view, manage, publish, block',
    'staff.invite, suspend, role.assign, session.revoke',
    'reports.view, export, configure; settings.manage; audit.view',
]: bullet(doc, item)
h(doc, '4.5 Emergency break-glass access', 2)
doc.add_paragraph('An authorized responder can temporarily view only the emergency card, critical allergies, medications, location, and contacts required for immediate care. The responder provides a reason; access expires; the event is logged and reviewable. Emergency access never exposes an entire medical history by default.')

h(doc, '5 Launch wedge Emergency Care Room')
doc.add_paragraph('The first memorable product is not a red SOS button. It is a shared Emergency Care Room connecting the patient, caregiver, dispatcher, ambulance crew, receiving hospital, doctor, and billing team around one incident.')
h(doc, '5.1 End to end flow', 2)
for item in [
    'Patient or caregiver selects the person needing help and starts an emergency request.',
    'A short safety check presents 112 guidance for severe symptoms.',
    'The request shares location and the minimum emergency card required for response.',
    'The network matches an ambulance by capability, readiness, distance, and service area.',
    'Dispatcher assigns a vehicle and crew; the caregiver sees status and ETA.',
    'The receiving hospital confirms emergency capacity and receives a pre-arrival summary.',
    'The doctor or emergency team records arrival and handoff events.',
    'Billing captures ambulance, consultation, and hospital charges against the incident.',
    'The incident closes into the patient timeline and an optional follow-up care plan.',
]: numbered(doc, item)
h(doc, '5.2 Differentiating pain-point features', 2)
add_table(doc, ['Feature', 'Pain removed', 'Demo proof'], [
    ('Capability-aware dispatch', 'Nearest vehicle may lack oxygen, ALS crew, or equipment', 'Vehicle cards show capability, readiness, ETA, and acceptance'),
    ('Hospital readiness handshake', 'Ambulance arrives without knowing capacity', 'Hospital console accepts pre-arrival and shows readiness'),
    ('Permission-aware incident room', 'Participants see too much or too little', 'Each workspace shows a safe view of one incident'),
    ('Escalation and fallback', 'A missed call costs critical time', 'Unaccepted request escalates to backup and calling guidance'),
    ('Handoff ledger', 'No accountability between ambulance and hospital', 'Timeline records people, times, status, and handoffs'),
    ('Post-event continuity', 'Care ends when ambulance arrives', 'Incident creates documents, billing, tasks, and follow-up'),
])

h(doc, '6 Provider operating system')
h(doc, '6.1 Hospital and clinic modules', 2)
for item in ['Organization onboarding and verification', 'Branch and department configuration', 'Provider profiles, credentials, fees, languages, and availability', 'Appointment slots, queues, check-in, wait-time updates, and no-show handling', 'Consent-aware patient registration and records', 'Invoices, receipts, refunds, discounts, settlements, and payment status', 'Staff invitation, role builder, branch/department scope, and access review', 'Operational dashboards for volume, wait time, capacity, cancellations, and revenue', 'Content and offer publishing for the patient app', 'Exportable reports and an audit log']:
    bullet(doc, item)
h(doc, '6.2 Doctor workspace', 2)
doc.add_paragraph('The doctor workspace focuses on the working day: schedule, queue, check-in status, patient context, reports, visit notes, prescriptions, follow-up tasks, and availability. It should reduce clicks during a consultation rather than imitate a large ERP.')
h(doc, '6.3 Independent practice support', 2)
doc.add_paragraph('An independent doctor can create a verified practice, configure a schedule, invite an assistant, publish services, accept appointments, issue invoices, and refer a patient to a hospital or diagnostic partner. ApnaCare can therefore create supply even before large hospital contracts.')

h(doc, '7 Patient experience')
for item in ['Emergency card with consent controls and family contacts', 'Care Circle for parents, children, and dependents', 'Provider discovery with transparent fit reasons, not only star ratings', 'Appointments, queues, reschedule, cancellation, and reminders', 'Health vault for reports, prescriptions, and discharge summaries', 'Care timeline linking emergencies, visits, bills, documents, and follow-ups', 'Low-data and Hindi/English-friendly emergency access', 'Care Companion for navigation, organization, and safe report explanation', 'Future medicine delivery and diagnostic fulfillment']:
    bullet(doc, item)

h(doc, '8 Care Companion AI')
doc.add_paragraph('The assistant is a tool-using care navigator. It creates value by taking actions inside ApnaCare, not by producing generic health paragraphs.')
add_table(doc, ['Capability', 'Allowed behavior', 'Safety boundary'], [
    ('Guided intake', 'Convert plain language into structured care needs', 'Never claim a diagnosis'),
    ('Emergency routing', 'Detect urgent language and open emergency actions', 'Show 112 guidance and escalate'),
    ('Provider matching', 'Filter by readiness, fee, language, distance, and slot', 'Explain the match using source data'),
    ('Booking agent', 'Create, reschedule, and cancel after confirmation', 'Never silently commit'),
    ('Report explainer', 'Summarize documents and draft doctor questions', 'Never prescribe or change treatment'),
    ('Care organizer', 'Create reminders, tasks, and caregiver updates', 'Respect consent and membership scope'),
])

h(doc, '9 Business model and go to market')
add_table(doc, ['Revenue stream', 'Buyer', 'Value'], [
    ('Provider subscription', 'Hospital, clinic, ambulance company', 'Seats, workflows, billing, reporting, branding, support'),
    ('Network transaction fee', 'Provider or payer', 'Qualified appointment and emergency demand'),
    ('Premium family plan', 'Patient or caregiver', 'Care Circle, storage, reminders, coordination'),
    ('Enterprise modules', 'Hospital network', 'Integrations, analytics, SLA, multi-branch controls'),
    ('Future pharmacy and diagnostics margin', 'Partner network', 'Convenient fulfillment after trust is established'),
])
doc.add_paragraph('Start with one city, one hospital or clinic partner, and a verified ambulance network. Emergency readiness creates patient demand; provider tools create retention and revenue. Appointments, records, bills, and medicines create recurring usage.')

h(doc, '10 Frontend implementation plan')
add_table(doc, ['Phase', 'Frontend outcome', 'Dummy scenario'], [
    ('1. Platform shell', 'Role-aware navigation, organization switcher, permission-aware actions', 'One doctor belongs to a clinic and ambulance provider'),
    ('2. Emergency journey', 'Patient SOS, dispatcher queue, ambulance status, hospital readiness, timeline', 'Sunita Jain requests cardiac support and reaches Apollo emergency'),
    ('3. Provider operations', 'Hospital/clinic dashboard, doctor day view, queues, services, billing', 'Apollo changes cardiology slots and issues an invoice'),
    ('4. Patient continuity', 'Care Circle, vault, tasks, reminders, and follow-up', 'The incident becomes a cardiology plan'),
    ('5. AI prototype', 'Care Companion with safe responses and action previews', 'Assistant finds a doctor or starts emergency flow'),
    ('6. Review checkpoint', 'All routes work from local repositories and testable state', 'Reviewer can explore every persona without backend'),
])
doc.add_paragraph('Keep the existing Flutter green visual language. New screens should use shared models, providers, and repository interfaces so a future backend replaces dummy data without rewriting presentation code.')

h(doc, '11 Backend-ready architecture')
doc.add_paragraph('The backend must support multi-tenancy, authentication, relational data, object storage, realtime events, notifications, and audit logging. The frontend can start with Supabase or an equivalent service, but its contracts should remain provider-agnostic.')
add_table(doc, ['Domain', 'Entities'], [
    ('Identity and access', 'Person, organization, branch, membership, role, permission, policy, session'),
    ('Care', 'Patient, care relationship, emergency card, consent, document, task, timeline event'),
    ('Providers', 'Doctor, clinic, hospital, ambulance provider, vehicle, capability, service, schedule'),
    ('Operations', 'Appointment, queue, encounter, handoff, emergency incident, notification'),
    ('Finance', 'Invoice, line item, payment, refund, settlement, subscription, commission'),
    ('Trust', 'Verification case, audit event, incident review, support case, SLA metric'),
])

h(doc, '12 Metrics')
add_table(doc, ['Area', 'Metrics'], [
    ('Emergency', 'Request-to-accept, accept-to-arrival, fallback rate, hospital acknowledgement'),
    ('Patient', 'Emergency-card completion, booking conversion, repeat family usage, no-show rate'),
    ('Provider', 'Weekly active staff, wait time, slot utilization, billing completion, time saved'),
    ('Network', 'Verified active providers, coverage by zone, successful handoffs, referrals'),
    ('Business', 'Provider activation, subscription retention, revenue per organization, transaction margin'),
    ('Trust', 'Unauthorized access, audit completion, support resolution, safety escalation'),
])

h(doc, '13 Risks and non-goals')
for item in ['Do not promise diagnosis, guaranteed ambulance ETA, or a guaranteed hospital bed.', 'Do not launch a huge ERP surface before proving emergency and appointment workflows.', 'Do not make the frontend the only security layer; backend policies and database isolation are mandatory.', 'Do not allow unrestricted custom roles without scope, templates, and audit.', 'Do not treat ratings as the only quality signal; show readiness, verification, wait time, and limitations.', 'Do not add medicine delivery until emergency coordination and provider operations are reliable.']:
    bullet(doc, item)

h(doc, '14 Build prompt for future implementation')
doc.add_paragraph('Build ApnaCare as a multi-tenant emergency-to-everyday healthcare platform. Preserve the existing Flutter visual language and use dummy repositories first. Implement a patient emergency flow, dispatcher workspace, ambulance status, hospital pre-arrival console, doctor day view, provider admin dashboard, billing snapshots, Care Circle records, and an auditable incident timeline. Model person identity separately from organizations, memberships, roles, permissions, scopes, policies, and emergency break-glass access. Provider owners receive full control of their tenant and can create custom roles from platform permission primitives. Keep every feature backend-ready behind interfaces and state providers. AI must be a safe tool-using Care Companion for navigation, bookings, document summaries, and reminders; it must not diagnose or prescribe. Provide review checkpoints after each meaningful feature batch and do not commit or push changes without explicit approval.')

h(doc, '15 Immediate decision')
doc.add_paragraph('The next implementation checkpoint is the Emergency Care Room with the role-aware access foundation. It is small enough to finish with dummy data, large enough to demonstrate the company, and directly connected to patient, ambulance, hospital, doctor, billing, and admin surfaces.')

for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('ApnaCare Product Requirements Document').font.size = Pt(8)

doc.save(OUT)
print(OUT)
